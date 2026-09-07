"""touchbistro.py
================
Consolidated TouchBistro reporting script.

Reports implemented
-------------------
  run_order_type()       —  Sales by Order Type  (cols H–M)
  run_sales_by_section() —  Sales by Section     (cols B–G)
  run_sales_by_category()—  Sales by Category    (cols N+, dynamic)
  run_sales_by_hour()    —  Sales by Hour        (cols under "Daily and Hourly Sales" header)

How it works
------------
1. Selenium starts a fresh Edge window and logs in automatically.
2. Every API call is made through the browser via execute_async_script so
   the browser's own Okta session handles auth — nothing to copy or manage.
3. Results are summed from daily records (the API has no summary row).
4. Data is written into the correct Excel columns and the workbook is saved.

Usage
-----
Dates are auto-computed from the current date (last completed Mon–Sun week).
To override, set "auto": false in config.json and fill in the date fields.

Run:  python touchbistro.py
"""

import base64
import json
import logging
import os
import sys
import tempfile
from copy import copy
from datetime import datetime, timezone, timedelta
from dotenv import load_dotenv
import openpyxl
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.edge.service import Service
from selenium.webdriver.edge.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))


# ════════════════════════════════════════════════════════════════
# CONFIG
# ════════════════════════════════════════════════════════════════

DRIVER_PATH = "msedgedriver.exe"
MASTER_XLSX = r"..\Weekly Reports\Weekly Sales summary (week 5).xlsx"
CHECKER_TEMPLATE_XLSX  = r"..\checker_template.xlsx"

# TouchBistro login credentials  (loaded from .env)
TB_USERNAME = os.environ["TB_USERNAME"]
TB_PASSWORD = os.environ["TB_PASSWORD"]


def _format_week_label(start, end):
    """Format a date range as 'Jul 6 - Jul 12' (no zero-padding)."""
    return f"{start.strftime('%b')} {start.day} - {end.strftime('%b')} {end.day}"


def compute_week_dates():
    """
    Compute the last completed Mon–Sun week from today's date.
    Returns (prev_week_label, target_week_label, week_start, week_end).
    """
    from datetime import date
    today = date.today()
    # days_since_sunday: Mon→1, Tue→2, … Sun→7
    days_since_sunday = (today.weekday() + 1) % 7 or 7
    last_sunday = today - timedelta(days=days_since_sunday)
    last_monday = last_sunday - timedelta(days=6)

    prev_sunday = last_sunday - timedelta(days=7)
    prev_monday = last_monday - timedelta(days=7)

    return (
        _format_week_label(prev_monday, prev_sunday),
        _format_week_label(last_monday, last_sunday),
        last_monday.strftime("%Y-%m-%d"),
        last_sunday.strftime("%Y-%m-%d"),
    )


def _load_config():
    """
    Determine week dates.
    If config.json exists and has "auto": false  →  use manual values.
    Otherwise  →  auto-compute from today's date.
    """
    config_path = os.path.join(os.path.dirname(__file__), "config.json")
    try:
        with open(config_path, "r") as f:
            cfg = json.load(f)
        if not cfg.get("auto", True):
            # Manual override mode
            return (
                cfg["prev_week_label"],
                cfg["target_week_label"],
                cfg["week_start"],
                cfg["week_end"],
            )
    except (FileNotFoundError, json.JSONDecodeError):
        pass  # Fall through to auto-compute

    return compute_week_dates()


PREV_WEEK_LABEL, TARGET_WEEK_LABEL, WEEK_START, WEEK_END = _load_config()

# Output path for the menu items workbook (uses TARGET_WEEK_LABEL as filename)
MENU_ITEMS_OUTPUT_XLSX = fr"..\Weekly Item Reports\{TARGET_WEEK_LABEL}.xlsx"


# ════════════════════════════════════════════════════════════════
# VENUE CONFIG
# ════════════════════════════════════════════════════════════════

LOGIN_URL = "https://login.touchbistro.com/login/login.htm"
ADMIN_URL = "https://admin.touchbistro.com"

# API URL templates  —  {venue_id}, {start}, {end} are filled at runtime
ORDER_TYPE_URL = (
    ADMIN_URL + "/api/frontend/report/v1/venues/{venue_id}/reports/order-type"
    + "?start={start}&end={end}"
)
SECTION_URL = (
    ADMIN_URL + "/api/frontend/report/v1/venues/{venue_id}/reports/sales-by-section"
    + "?start={start}&end={end}"
)
CATEGORY_URL = (
    ADMIN_URL + "/api/frontend/report/v1/venues/{venue_id}/reports/sales-by-sales-category"
    + "?start={start}&end={end}&orderType=any"
)
HOURLY_URL = (
    ADMIN_URL + "/api/frontend/report/v1/venues/{venue_id}/reports/sales-hourly-net-bill-start"
    + "?start={start}&end={end}"
)
MENU_ITEM_URL = (
    ADMIN_URL + "/api/frontend/report/v1/venues/{venue_id}/reports/export_data"
    + "?start={start}&end={end}&report_name=sales_by_menu_item&type=xlsx"
)

# venue_id → (display name, Excel sheet name)
VENUES = {
    56776: ("Karachi Kabab Wala - Queen Street", "Kababwala - Queen"),
    55119: ("Pizza Karachi- Eglinton",           "Pizza K Eglinton"),
    55118: ("Pizza Karachi -Heartland",          "Pizza K Heartland"),
    51879: ("Karachi Kabab Wala",                "Kababwala"),
    51876: ("Karachi Food Court",                "Karachi Food Court"),
    54708: ("Pizza Karachi Downtown TO",         "Queen St."),
    52043: ("Pizza Karachi- Highway Karahi",     "Highway"),
    51880: ("Pizza Karachi - Wonderland",        "Jane"),
    51878: ("Pizza Karachi - Lebovic",           "Lebovic"),
    51877: ("Pizza Karachi - Ajax",              "Ajax"),
    51594: ("Pizza Karachi - Markham Rd",        "Markham"),
    65879: ("Pizza Karachi - Oshawa",            "Oshawa"),
    66035: ("Pizza Karachi - Milton",            "Milton")
}

# venue_id → sheet name as expected by the checker's SUMIF formulas
MENU_ITEM_SHEET_NAMES = {
    56776: "Karachi Kabab Wala - Queen",
    55119: "Eglinton",
    55118: "Heartland",
    51879: "Karachi Kabab Wala",
    51876: "Karachi Food Court",
    54708: "Pizza Karachi Downtown TO",
    52043: "Highway Karahi",
    51880: "Wonderland",
    51878: "Lebovic",
    51877: "Ajax",
    51594: "Markham Rd",
    65879: "Oshawa",
    66035: "Milton"
}

# Time buckets for the hourly report (label, [hour strings from API])
HOUR_BUCKETS = [
    ("8AM - 11AM",  ["8am - 9am", "9am - 10am", "10am - 11am"]),
    ("11AM - 3PM",  ["11am - 12pm", "12pm - 1pm", "1pm - 2pm", "2pm - 3pm"]),
    ("3PM - 6PM",   ["3pm - 4pm", "4pm - 5pm", "5pm - 6pm"]),
    ("6PM - 11PM",  ["6pm - 7pm", "7pm - 8pm", "8pm - 9pm", "9pm - 10pm", "10pm - 11pm"]),
    ("11PM - 7AM",  ["11pm - 12am", "12am - 1am", "1am - 2am", "2am - 3am",
                     "3am - 4am", "4am - 5am", "5am - 6am", "6am - 7am"]),
]

# Delivery platform section names as they appear in the API response
DELIVERY_PLATFORMS = {
    "skip the dishes": "skip",
    "doordash":        "dd",
    "ubereats":        "uber",
}

# Excel column constants (1-based)
SECTION_BILL_COUNT_COL  = 2   # B
SECTION_NET_SALES_COL   = 3   # C
SECTION_REMAINING_COL   = 4   # D
SECTION_SKIP_COL        = 5   # E
SECTION_DD_COL          = 6   # F
SECTION_UBER_COL        = 7   # G
ORDER_TYPE_START_COL    = 8   # H  (Take Out … Net Sales → H–M)


# ════════════════════════════════════════════════════════════════
# HELPERS — general
# ════════════════════════════════════════════════════════════════

class _TeeWriter:
    """Write to multiple streams simultaneously (for logging to file + console)."""
    def __init__(self, *streams):
        self.streams = streams
    def write(self, data):
        for s in self.streams:
            s.write(data)
            s.flush()
    def flush(self):
        for s in self.streams:
            s.flush()

def to_unix(date_str: str, add_day: bool = False) -> int:
    """Convert 'YYYY-MM-DD' to a Unix timestamp at midnight UTC."""
    dt = datetime.strptime(date_str, "%Y-%m-%d").replace(tzinfo=timezone.utc)
    if add_day:
        dt += timedelta(days=1)
    return int(dt.timestamp())


def login() -> webdriver.Edge:
    """Start a fresh Edge window and log in to admin.touchbistro.com."""
    opts = Options()
    opts.use_chromium = True
    driver = webdriver.Edge(service=Service(DRIVER_PATH), options=opts)
    wait = WebDriverWait(driver, 30)

    print("Opening login page...")
    driver.get(LOGIN_URL)

    wait.until(EC.presence_of_element_located((By.ID, "okta-signin-username"))).send_keys(TB_USERNAME)
    driver.find_element(By.ID, "okta-signin-password").send_keys(TB_PASSWORD)
    driver.find_element(By.ID, "okta-signin-submit").click()

    print("Waiting for login to complete...")
    wait.until(EC.url_contains("admin.touchbistro.com"))
    print("✅ Logged in.\n")
    return driver


def _log_api_response(label: str, url: str, payload):
    """Write the captured response to stdout so it is also stored in the log file."""
    try:
        formatted = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True)
    except TypeError:
        formatted = str(payload)
    print(f"{label} for {url}:\n{formatted}")


def fetch_via_browser(driver, url: str) -> list[dict]:
    """
    Make a GET request through the browser so its Okta session is used.
    Returns parsed JSON or raises RuntimeError on HTTP/network error.
    """
    result = driver.execute_async_script("""
        var url      = arguments[0];
        var callback = arguments[1];
        fetch(url)
            .then(function(r) {
                if (!r.ok) {
                    callback({ __error__: r.status + ' ' + r.statusText });
                } else {
                    r.json().then(function(d) { callback(d); });
                }
            })
            .catch(function(e) { callback({ __error__: e.toString() }); });
    """, url)

    if isinstance(result, dict) and "__error__" in result:
        print(f"API fetch failed for {url}: {result['__error__']}")
        raise RuntimeError(f"API error: {result['__error__']}")

    _log_api_response("API response", url, result)
    return result


def find_target_row(ws) -> int:
    """
    Return the row to write into.
    Uses TARGET_WEEK_LABEL if it exists; otherwise creates a row below PREV_WEEK_LABEL.
    """
    prev_row = None
    for cell in ws["A"]:
        val = str(cell.value).strip() if cell.value is not None else ""
        if val == TARGET_WEEK_LABEL:
            return cell.row
        if val == PREV_WEEK_LABEL:
            prev_row = cell.row

    if prev_row is None:
        raise RuntimeError(
            f"Neither '{TARGET_WEEK_LABEL}' nor '{PREV_WEEK_LABEL}' found "
            f"in column A of sheet '{ws.title}'. Check CONFIG labels."
        )
    target_row = prev_row + 1
    ws.cell(row=target_row, column=1, value=TARGET_WEEK_LABEL)
    print(f"   ℹ️  Created row {target_row} for '{TARGET_WEEK_LABEL}'.")
    return target_row


def build_url(template: str, venue_id: int, start_ts: int, end_ts: int) -> str:
    return template.format(venue_id=venue_id, start=start_ts, end=end_ts)


# ════════════════════════════════════════════════════════════════
# REPORT 1 — Sales by Order Type
# ════════════════════════════════════════════════════════════════

def run_order_type(driver, wb, start_ts, end_ts):
    """
    Sums daily order-type records and writes to cols H–M:
      Take Out | Delivery | Bar Tab | Dine-In | Online Order | Net Sales
    """
    print("── Sales by Order Type ──────────────────────────")
    keys = ["takeout", "delivery", "bartab", "dinein", "onlineorder", "total"]
    col_labels = ["Take Out", "Delivery", "Bar Tab", "Dine-In", "Online Order", "Net Sales"]

    for venue_id, (name, sheet_name) in VENUES.items():
        print(f"▶  {name}")
        url = build_url(ORDER_TYPE_URL, venue_id, start_ts, end_ts)
        try:
            records = fetch_via_browser(driver, url)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        if not records:
            print("   ⚠️  No data — skipping."); continue

        values = [round(sum(float(r[k]) for r in records), 2) for k in keys]
        for lbl, val in zip(col_labels, values):
            print(f"   {lbl:<14}: {val:>10.2f}")

        ws = wb[sheet_name]
        row = find_target_row(ws)
        for idx, val in enumerate(values):
            ws.cell(row=row, column=ORDER_TYPE_START_COL + idx, value=val)
        print(f"   ✅ Written to row {row}")


# ════════════════════════════════════════════════════════════════
# REPORT 2 — Sales by Section
# ════════════════════════════════════════════════════════════════

def run_sales_by_section(driver, wb, start_ts, end_ts):
    """
    Classifies sections into delivery platforms vs. in-house and writes to cols B–G:
      Total Bill Count | Net Sales | Remaining | Skip | DoorDash | Uber
    """
    print("\n── Sales by Section ─────────────────────────────")

    for venue_id, (name, sheet_name) in VENUES.items():
        print(f"▶  {name}")
        url = build_url(SECTION_URL, venue_id, start_ts, end_ts)
        try:
            records = fetch_via_browser(driver, url)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        if not records:
            print("   ⚠️  No data — skipping."); continue

        total_bill_count = sum(r["bill_count"] for r in records)
        net_sales = skip = dd = uber = remaining = 0.0

        for r in records:
            rev  = float(r["sales_revenue"])
            name_lower = r["section_name"].strip().lower()
            net_sales += rev
            if name_lower == "skip the dishes":
                skip += rev
            elif name_lower == "doordash":
                dd += rev
            elif name_lower == "ubereats":
                uber += rev
            else:
                remaining += rev

        print(f"   {'Bill Count':<14}: {total_bill_count:>10}")
        print(f"   {'Net Sales':<14}: {net_sales:>10.2f}")
        print(f"   {'Remaining':<14}: {remaining:>10.2f}")
        print(f"   {'Skip':<14}: {skip:>10.2f}")
        print(f"   {'DoorDash':<14}: {dd:>10.2f}")
        print(f"   {'Uber':<14}: {uber:>10.2f}")

        ws = wb[sheet_name]
        row = find_target_row(ws)
        cell = ws.cell(row=row, column=SECTION_BILL_COUNT_COL, value=int(total_bill_count))
        cell.number_format = "0"
        ws.cell(row=row, column=SECTION_NET_SALES_COL,  value=round(net_sales, 2))
        ws.cell(row=row, column=SECTION_REMAINING_COL,  value=round(remaining, 2))
        ws.cell(row=row, column=SECTION_SKIP_COL,       value=round(skip, 2))
        ws.cell(row=row, column=SECTION_DD_COL,         value=round(dd, 2))
        ws.cell(row=row, column=SECTION_UBER_COL,       value=round(uber, 2))
        print(f"   ✅ Written to row {row}")


# ════════════════════════════════════════════════════════════════
# REPORT 3 — Sales by Category
# ════════════════════════════════════════════════════════════════

def _build_category_header_map(ws) -> dict[str, int]:
    """
    Scan row 2 of the sheet for column headers.
    Returns {header_lower: column_index} for every non-empty header cell.
    """
    header_map = {}
    for cell in ws[2]:
        if cell.value is not None:
            header_map[str(cell.value).strip().lower()] = cell.column
    return header_map


def _match_category(category_name: str, header_map: dict[str, int]) -> int | None:
    """
    Find the column for a category by checking if any Excel header
    starts with the category name (case-insensitive), e.g. header
    "Food (Pizza K)" matches API category "Food".
    """
    cat_lower = category_name.strip().lower()
    for header, col in header_map.items():
        if header.startswith(cat_lower):
            return col
    return None


def _find_category_net_sales_col(ws) -> int | None:
    """
    Return the column index of the 3rd 'Net Sales' header in row 2.
    The first two belong to Sales by Section and Sales by Order Type;
    the third is the category total column.
    """
    count = 0
    for cell in ws[2]:
        if cell.value is not None and str(cell.value).strip() == "Net Sales":
            count += 1
            if count == 3:
                return cell.column
    return None


def _insert_category_column(ws, at_col: int, header: str):
    """
    Insert a new column at at_col, write the category name as the row-2
    header, and correctly update every row-1 merged header.

    We take full ownership of all row-1 merges:
      1. Snapshot every merge in row 1 before the insert.
      2. Call insert_cols (shifts cell data correctly).
      3. Discard ALL row-1 merges from the set (bypassing unmerge_cells to
         avoid its KeyError when cells have already been shifted).
      4. Re-apply each merge with boundaries computed from the pre-insert
         snapshot — category merge extends by 1, all others shift normally.
    """
    # 1. Snapshot all row-1 merges (min_col, max_col, title) BEFORE insert
    row1_merges = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row == 1 and m.max_row == 1:
            row1_merges.append({
                "min_col": m.min_col,
                "max_col": m.max_col,
                "title":   ws.cell(1, m.min_col).value,
            })

    # 2. Insert the new column and write its row-2 header
    ws.insert_cols(at_col)
    ws.cell(row=2, column=at_col, value=header)

    # 3. Discard every row-1 merge that openpyxl now has recorded
    #    (their ranges may be in an inconsistent state after the insert)
    for m in list(ws.merged_cells.ranges):
        if m.min_row == 1 and m.max_row == 1:
            ws.merged_cells.ranges.discard(m)

    # 4. Re-apply all row-1 merges with manually computed boundaries
    for info in row1_merges:
        old_min = info["min_col"]
        old_max = info["max_col"]
        title   = info["title"]
        is_cat  = title and "category" in str(title).lower()

        if is_cat and old_min <= at_col <= old_max + 1:
            # Category merge: extend right by 1 to cover the new column
            new_min, new_max = old_min, old_max + 1
        elif old_min >= at_col:
            # Entirely to the right of the insert: shift both ends right
            new_min, new_max = old_min + 1, old_max + 1
        elif old_max >= at_col:
            # Spans the insert point: expand the right end
            new_min, new_max = old_min, old_max + 1
        else:
            # Entirely to the left: no change
            new_min, new_max = old_min, old_max

        ws.merge_cells(start_row=1, start_column=new_min,
                       end_row=1,   end_column=new_max)
        ws.cell(row=1, column=new_min).value = title


def run_sales_by_category(driver, wb, start_ts, end_ts):
    """
    Writes each category's net_revenue to its matching Excel column and
    writes the grand total to the 3rd 'Net Sales' column (category total).

    If a category from the API is missing from the sheet, a new column is
    inserted immediately before the 'Net Sales' column, with the category
    name written as the row-2 header.
    """
    print("\n── Sales by Category ────────────────────────────")

    for venue_id, (name, sheet_name) in VENUES.items():
        print(f"▶  {name}")
        url = build_url(CATEGORY_URL, venue_id, start_ts, end_ts)
        try:
            records = fetch_via_browser(driver, url)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        if not records:
            print("   ⚠️  No data — skipping."); continue

        ws = wb[sheet_name]
        row = find_target_row(ws)
        total_net_revenue = 0.0

        for r in records:
            cat_name = r["sales_category_name"]
            net_rev  = round(float(r["net_revenue"]), 2)
            total_net_revenue += net_rev

            # Rebuild header map each iteration — it may shift after an insert
            header_map = _build_category_header_map(ws)
            col = _match_category(cat_name, header_map)

            if col is None:
                # Insert a new column before the category Net Sales column
                # and extend the merged header to keep Net Sales inside the group
                net_sales_col = _find_category_net_sales_col(ws)
                if net_sales_col is not None:
                    _insert_category_column(ws, at_col=net_sales_col, header=cat_name)
                    col = net_sales_col
                    print(f"   ➕ '{cat_name}' not in sheet — added column {col}, merge extended.")
                else:
                    print(f"   ⚠️  '{cat_name}' not found and no Net Sales col to insert before — skipping.")
                    continue

            ws.cell(row=row, column=col, value=net_rev)
            print(f"   {cat_name:<20}: {net_rev:>10.2f}  → col {col}")

        # Write the category grand total to the 3rd Net Sales column
        net_sales_col = _find_category_net_sales_col(ws)
        if net_sales_col is not None:
            ws.cell(row=row, column=net_sales_col, value=round(total_net_revenue, 2))
            print(f"   {'Net Sales (total)':<20}: {round(total_net_revenue, 2):>10.2f}  → col {net_sales_col}")
        else:
            print("   ⚠️  No 3rd 'Net Sales' column found — category total not written.")

        print(f"   ✅ Written to row {row}")


# ════════════════════════════════════════════════════════════════
# REPORT 4 — Sales by Hour
# ════════════════════════════════════════════════════════════════

_DAY_KEYS = [
    "sunday_net_sales", "monday_net_sales", "tuesday_net_sales",
    "wednesday_net_sales", "thursday_net_sales", "friday_net_sales",
    "saturday_net_sales",
]


def _find_hourly_start_col(ws) -> int:
    """
    Find the starting column of the 'Daily and Hourly Sales (TBD)'
    merged header in row 1. Raises if not found.
    """
    for m in ws.merged_cells.ranges:
        if m.min_row == 1:
            cell_val = ws.cell(m.min_row, m.min_col).value
            if cell_val and "hourly" in str(cell_val).lower():
                return m.min_col
    raise RuntimeError(
        f"'Daily and Hourly Sales' header not found in row 1 of sheet '{ws.title}'."
    )


def run_sales_by_hour(driver, wb, start_ts, end_ts):
    """
    Sums all 7 day columns for each hour slot, buckets them into 5 time
    ranges, and writes to the columns under the 'Daily and Hourly Sales'
    merged header.
    """
    print("\n── Sales by Hour ────────────────────────────────")

    for venue_id, (name, sheet_name) in VENUES.items():
        print(f"▶  {name}")
        url = build_url(HOURLY_URL, venue_id, start_ts, end_ts)
        try:
            records = fetch_via_browser(driver, url)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        if not records:
            print("   ⚠️  No data — skipping."); continue

        # Sum all 7 day columns per hour to get the week total for that slot
        hour_sums: dict[str, float] = {}
        for r in records:
            hour = r["hour_of_the_day"].strip().lower()
            hour_sums[hour] = sum(float(r[d]) for d in _DAY_KEYS)

        # Aggregate into time buckets
        bucket_values = []
        for bucket_label, hours in HOUR_BUCKETS:
            total = 0.0
            for h in hours:
                if h not in hour_sums:
                    print(f"   ℹ️  Hour '{h}' missing for {sheet_name}, treating as 0.")
                total += hour_sums.get(h, 0.0)
            bucket_values.append(round(total, 2))
            print(f"   {bucket_label:<14}: {bucket_values[-1]:>10.2f}")

        ws = wb[sheet_name]
        row = find_target_row(ws)

        try:
            start_col = _find_hourly_start_col(ws)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        for idx, val in enumerate(bucket_values):
            ws.cell(row=row, column=start_col + idx, value=val)
        print(f"   ✅ Written to row {row} starting at col {start_col}")


# ════════════════════════════════════════════════════════════════
# REPORT 5 — Sales by Menu Item
# ════════════════════════════════════════════════════════════════

def fetch_binary_via_browser(driver, url: str) -> bytes:
    """
    Fetch a binary file through the browser so its Okta session is used.
    Returns raw bytes (e.g. an XLSX file) or raises RuntimeError.
    """
    b64 = driver.execute_async_script("""
        var url      = arguments[0];
        var callback = arguments[1];
        fetch(url)
            .then(function(r) {
                if (!r.ok) {
                    callback({ __error__: r.status + ' ' + r.statusText });
                    return;
                }
                r.arrayBuffer().then(function(buf) {
                    var bytes  = new Uint8Array(buf);
                    var chunks = [];
                    for (var i = 0; i < bytes.length; i += 8192) {
                        chunks.push(String.fromCharCode.apply(null, bytes.subarray(i, i + 8192)));
                    }
                    callback(btoa(chunks.join('')));
                });
            })
            .catch(function(e) { callback({ __error__: e.toString() }); });
    """, url)

    if isinstance(b64, dict) and "__error__" in b64:
        print(f"Binary API fetch failed for {url}: {b64['__error__']}")
        raise RuntimeError(f"API error: {b64['__error__']}")

    raw = base64.b64decode(b64)
    print(f"Binary API response for {url}: {len(raw)} bytes (Excel export captured)")
    return raw


def _copy_checker_sheet(template_path: str, target_wb, week_label: str):
    """
    Copy the 'checker' sheet from the template file into target_wb,
    preserving all cell values, styles, merged cells, column widths,
    and row heights. Updates the week label in A1.
    """
    tmpl_wb = load_workbook(template_path, data_only=False)
    src = tmpl_wb["checker"]
    dst = target_wb.create_sheet(title="checker")

    # Values + styles
    for row in src.iter_rows():
        for cell in row:
            dst_cell = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                dst_cell.font         = copy(cell.font)
                dst_cell.fill         = copy(cell.fill)
                dst_cell.alignment    = copy(cell.alignment)
                dst_cell.border       = copy(cell.border)
                dst_cell.number_format = cell.number_format
                dst_cell.protection   = copy(cell.protection)

    # Merged cells
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))

    # Column widths
    for col, dim in src.column_dimensions.items():
        if dim.width:
            dst.column_dimensions[col].width = dim.width

    # Row heights
    for row_idx, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[row_idx].height = dim.height

    # Update week label (A1 is inside the A1:X2 merge, so write to top-left only)
    dst.cell(row=1, column=1, value=week_label)


def run_menu_items(driver, start_ts: int, end_ts: int):
    """
    Downloads the Sales by Menu Item XLSX for each venue via the browser,
    assembles them all into a single workbook, appends the checker sheet
    (copied from CHECKER_TEMPLATE_XLSX), and saves to MENU_ITEMS_OUTPUT_XLSX.
    """
    print("\n── Sales by Menu Item ───────────────────────────")

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)   # drop the default empty sheet
    # Force Excel to recalculate all formulas on open (checker SUMIF formulas)
    out_wb.calculation.calcMode      = "auto"
    out_wb.calculation.fullCalcOnLoad = True

    for venue_id, sheet_name in MENU_ITEM_SHEET_NAMES.items():
        display_name = VENUES[venue_id][0]
        print(f"▶  {display_name}")
        url = build_url(MENU_ITEM_URL, venue_id, start_ts, end_ts)

        try:
            raw = fetch_binary_via_browser(driver, url)
        except RuntimeError as exc:
            print(f"   ⚠️  {exc} — skipping."); continue

        # Write bytes to a temp file so openpyxl can open it
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
            tmp.write(raw)

        try:
            dl_wb = load_workbook(tmp_path, data_only=False)
            dl_ws = dl_wb.active
            target_ws = out_wb.create_sheet(title=sheet_name)

            for row in dl_ws.iter_rows():
                for cell in row:
                    val = cell.value
                    # Numeric columns (4–15) in data rows arrive as strings from TB
                    if cell.row >= 3 and 4 <= cell.column <= 15 and isinstance(val, str):
                        try:
                            clean = val.replace(",", "").strip()
                            val = float(clean.rstrip("%")) / 100 if clean.endswith("%") else float(clean)
                        except ValueError:
                            pass
                    target_ws.cell(row=cell.row, column=cell.column, value=val)

            print(f"   ✅ {dl_ws.max_row - 2} items → sheet '{sheet_name}'")
        finally:
            os.unlink(tmp_path)

    # Append the checker sheet from the template
    try:
        _copy_checker_sheet(CHECKER_TEMPLATE_XLSX, out_wb, TARGET_WEEK_LABEL)
        print("   ✅ Checker sheet copied from template.")
    except Exception as exc:
        print(f"   ⚠️  Could not copy checker: {exc}")

    os.makedirs(os.path.dirname(os.path.abspath(MENU_ITEMS_OUTPUT_XLSX)), exist_ok=True)
    out_wb.save(MENU_ITEMS_OUTPUT_XLSX)
    print(f"\n✅  Menu items workbook saved → {MENU_ITEMS_OUTPUT_XLSX}")


# ════════════════════════════════════════════════════════════════
# MAIN
# ════════════════════════════════════════════════════════════════

def main():
    # Set up logging — tee stdout to a log file
    log_dir = os.path.join(os.path.dirname(__file__), "..", "logs")
    os.makedirs(log_dir, exist_ok=True)
    log_file_path = os.path.join(
        log_dir,
        f"touchbistro_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.log",
    )
    log_file = open(log_file_path, "w", encoding="utf-8")
    sys.stdout = _TeeWriter(sys.__stdout__, log_file)

    print(f"Week : {WEEK_START} → {WEEK_END}")
    print(f"Label: '{TARGET_WEEK_LABEL}'\n")

    start_ts = to_unix(WEEK_START)
    end_ts   = to_unix(WEEK_END, add_day=True)
    print(f"Timestamps: start={start_ts}  end={end_ts}\n")

    driver = login()

    try:
        wb = load_workbook(MASTER_XLSX)

        run_order_type(driver, wb, start_ts, end_ts)
        run_sales_by_section(driver, wb, start_ts, end_ts)
        run_sales_by_category(driver, wb, start_ts, end_ts)
        run_sales_by_hour(driver, wb, start_ts, end_ts)

        wb.save(MASTER_XLSX)
        print(f"\n✅  Weekly summary done. Workbook saved → {MASTER_XLSX}")

        run_menu_items(driver, start_ts, end_ts)

    finally:
        driver.quit()
        sys.stdout = sys.__stdout__
        log_file.close()
        print(f"Log saved → {log_file_path}")


if __name__ == "__main__":
    main()
